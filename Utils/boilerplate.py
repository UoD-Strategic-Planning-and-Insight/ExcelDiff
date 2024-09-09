"""
This package contains functions that reduce common sets of instructions used in scripts down to single function calls.
"""

import openpyxl
from openpyxl.workbook import Workbook
from openpyxl.worksheet.table import Table
from openpyxl.worksheet.worksheet import Worksheet

from Utils.tables import TableUtil


def load_table(xlsx_path: str, sheet_name: str, table_name: str) -> TableUtil:
    """
    Opens a table with the given name, in a sheet with the given name, in an Excel file at the given path.
    :return: A TableUtil object wrapping the OpenPyxl table specified by the given path, sheet name, and table name.
    """

    wb: Workbook = openpyxl.load_workbook(xlsx_path,
                                          data_only=True)

    ws: Worksheet = wb[sheet_name]
    tbl: Table = ws.tables[table_name]
    tbl_util: TableUtil = TableUtil(xlsx_path, wb, ws, tbl)

    return tbl_util
